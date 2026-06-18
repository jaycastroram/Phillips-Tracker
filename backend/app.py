from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
import secrets
import sqlite3
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any, Literal

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, Response, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, EmailStr, Field

try:
    import cloudinary
    import cloudinary.uploader
except ImportError:  # pragma: no cover - uploads are disabled until the dependency is installed.
    cloudinary = None

try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError:  # pragma: no cover - local SQLite dev can run without psycopg installed.
    psycopg = None
    dict_row = None


BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "tracker.db"
SEED_ITEMS_PATH = BASE_DIR / "data" / "seed_items.json"
FRONTEND_DIST = BASE_DIR.parent / "frontend" / "dist"


def load_local_env() -> None:
    env_path = BASE_DIR / ".env"
    if not env_path.exists():
        return

    for line in env_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        os.environ.setdefault(key.strip().lstrip("\ufeff"), value.strip().strip('"').strip("'"))


load_local_env()

DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
IS_POSTGRES = DATABASE_URL.startswith(("postgres://", "postgresql://"))
SESSION_DAYS = 7
ADMIN_EMAIL = os.getenv("ADMIN_EMAIL", "admin@example.com")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "Admin123")
ADMIN_NAME = os.getenv("ADMIN_NAME", "Dev Admin")
CLOUDINARY_CLOUD_NAME = os.getenv("CLOUDINARY_CLOUD_NAME", "").strip()
CLOUDINARY_API_KEY = os.getenv("CLOUDINARY_API_KEY", "").strip()
CLOUDINARY_API_SECRET = os.getenv("CLOUDINARY_API_SECRET", "").strip()
CORS_ORIGINS = [
    origin.strip()
    for origin in os.getenv(
        "CORS_ORIGINS",
        "http://localhost:5173,http://127.0.0.1:5173",
    ).split(",")
    if origin.strip()
]
DB_INTEGRITY_ERRORS = (sqlite3.IntegrityError,)
if psycopg is not None:
    DB_INTEGRITY_ERRORS = (sqlite3.IntegrityError, psycopg.IntegrityError)

if cloudinary is not None:
    cloudinary.config(
        cloud_name=CLOUDINARY_CLOUD_NAME,
        api_key=CLOUDINARY_API_KEY,
        api_secret=CLOUDINARY_API_SECRET,
        secure=True,
    )

SHEETS = {
    "ad-hoc": {
        "label": "Ad Hoc",
        "dateLabel": "Request Date",
        "statuses": [
            "Details Needed",
            "Quoting",
            "Pending Feedback",
            "Creative",
            "Processing OA",
            "OA Pending Approval",
            "PPS Underway",
            "PPS In Transit",
            "PPS Pending Approval",
            "Order In Production",
            "Order In Transit",
            "Delivered",
            "ON HOLD",
        ],
    },
    "buys": {
        "label": "Buys",
        "dateLabel": "Buy",
        "statuses": [
            "Details Needed",
            "Quoting",
            "Pending Feedback",
            "Creative",
            "FINAL Quoting",
            "Canceled",
            "Processing OA",
            "OA Pending Approval",
            "PPS Underway",
            "PPS In Transit",
            "PPS Pending Approval",
            "Order In Production",
            "Order In Transit",
            "Delivered",
            "ON HOLD",
        ],
    },
    "completed": {
        "label": "Completed",
        "dateLabel": "Date/Buy",
        "statuses": [
            "Details Needed",
            "Quoting",
            "Pending Feedback",
            "Creative",
            "Processing OA",
            "OA Pending Approval",
            "PPS Underway",
            "PPS In Transit",
            "PPS Pending Approval",
            "Order In Production",
            "Order In Transit",
            "Delivered",
            "ON HOLD",
        ],
    },
}

EDITABLE_FIELDS = {
    "date_or_buy",
    "current_status",
    "visual_reference",
    "brand",
    "program_name",
    "item_name",
    "qty",
    "important_notes",
    "mrl_order_number",
    "estimated_ship_date",
    "estimated_ihd",
    "tracking",
}
ALLOWED_UPLOAD_TYPES = {"image/jpeg", "image/png", "image/webp"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
DEFAULT_KANBAN_COLUMNS = [
    {
        "title": "Needs Details",
        "statuses": ["Details Needed"],
        "sort_order": 1,
    },
    {
        "title": "Quoting / Feedback",
        "statuses": ["Quoting", "Pending Feedback", "FINAL Quoting"],
        "sort_order": 2,
    },
    {
        "title": "Creative / OA",
        "statuses": ["Creative", "Processing OA", "OA Pending Approval"],
        "sort_order": 3,
    },
    {
        "title": "PPS",
        "statuses": ["PPS Underway", "PPS In Transit", "PPS Pending Approval"],
        "sort_order": 4,
    },
    {
        "title": "Production / Transit",
        "statuses": ["Order In Production", "Order In Transit"],
        "sort_order": 5,
    },
    {
        "title": "On Hold / Canceled",
        "statuses": ["ON HOLD", "Canceled"],
        "sort_order": 6,
    },
]
DEFAULT_SURVEY_ITEMS = [
    {
        "item_name": "Tito's Bar Mat",
        "brand": "Tito's",
        "channel": "On Premise",
        "product_type": "Barware",
        "key_program": "Evergreen",
        "item_description": "Custom 16\" x 16\" black PVC service bar mat featuring custom molded oval nibs and Tito's branding.",
        "uom": "1/EA",
        "price": "$7.45",
        "image_url": "https://cincoro.sb.myfsionline.com/images/4006_10001_1.png",
        "sort_order": 1
    },
    {
        "item_name": "Tito's Round Bottle Presenter",
        "brand": "Tito's",
        "channel": "On Premise",
        "product_type": "Displays",
        "key_program": "Evergreen",
        "item_description": "Illuminated metal bottle presenter featuring laser-cut Tito's branding, integrated LED lighting, and illuminated side panels. Designed to fit 750ml, 1L, and 1.75L Tito's bottles",
        "uom": "1/EA",
        "price": "$332.04",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10002_1.png",
        "sort_order": 2
    },
    {
        "item_name": "Tito's Crop Top",
        "brand": "Tito's",
        "channel": "Off Premise",
        "product_type": "Wearables",
        "key_program": "Evergreen",
        "item_description": "Black athletic tank made from a nylon-spandex performance blend, featuring a Tito's logo chest imprint.",
        "uom": "1/EA",
        "price": "$19.75",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10003_1.png",
        "sort_order": 3
    },
    {
        "item_name": "Tito's Short Denim Romper",
        "brand": "Tito's",
        "channel": "Off Premise",
        "product_type": "Wearables",
        "key_program": "Evergreen",
        "item_description": "Custom cream denim romper featuring a Tito's logo sleeve imprint",
        "uom": "1/EA",
        "price": "$36.56",
        "image_url": "https://phillips.pdccatalog.com/images/3185_35838_1.png",
        "sort_order": 4
    },
    {
        "item_name": "Tito's Doghouse Display",
        "brand": "Tito's",
        "channel": "Off Premise",
        "product_type": "Displays",
        "key_program": "Evergreen",
        "item_description": "Custom wood doghouse display with a functional chalkboard roof, Tito's branded graphics, and a built-in center divider",
        "uom": "1/EA",
        "price": "$210.82",
        "image_url": "https://demo.mrlsmartbuy.com/images/3185_35838-BlueGraphics_1.png",
        "sort_order": 5
    },
    {
        "item_name": "Tito's Acrylic Top Garnish Caddy",
        "brand": "Tito's",
        "channel": "On Premise",
        "product_type": "Barware",
        "key_program": "Evergreen",
        "item_description": "Custom garnish caddy with Tito's branding, removable food-safe garnish trays, a clear lid, and copper accent detailing.",
        "uom": "1/EA",
        "price": "$41.72",
        "image_url": "https://phillips.pdccatalog.com/images/3185_38404_1.png",
        "sort_order": 6
    },
    {
        "item_name": "Tito's Napkin Caddy",
        "brand": "Tito's",
        "channel": "On Premise",
        "product_type": "Barware",
        "key_program": "Evergreen",
        "item_description": "Wood napkin caddy featuring Tito's branding, copper accent detailing, and a protective clear-coated finish for bar and tabletop use.",
        "uom": "1/EA",
        "price": "$19.59",
        "image_url": "https://cincoro.sb.myfsionline.com/images/4006_10001_1.png",
        "sort_order": 7
    },
    {
        "item_name": "Tito's Swimsuit",
        "brand": "Tito's",
        "channel": "Off Premise",
        "product_type": "Wearables",
        "key_program": "Evergreen",
        "item_description": "Custom orange swimsuit featuring Tito's branded graphics and a white belt.",
        "uom": "1/EA",
        "price": "$29.76",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10002_1.png",
        "sort_order": 8
    },
    {
        "item_name": "LALO Sales Booklet",
        "brand": "LALO",
        "channel": "Off Premise",
        "product_type": "Print",
        "key_program": "Evergreen",
        "item_description": "Folded Sales Book featuring imprint on all pages. Updated Lalo 2026 legal line.",
        "uom": "25/PK",
        "price": "$19.75",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10003_1.png",
        "sort_order": 9
    },
    {
        "item_name": "LALO Hat",
        "brand": "LALO",
        "channel": "Giveaways",
        "product_type": "Wearables",
        "key_program": "Sweeps",
        "item_description": "Custom Lalo branded hat with 3D embroidery. Includes tear away white labels.",
        "uom": "1/EA",
        "price": "$6.59",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10004_1.png",
        "sort_order": 10
    },
    {
        "item_name": "LALO Rail Mat",
        "brand": "LALO",
        "channel": "On Premise",
        "product_type": "Barware",
        "key_program": "Evergreen",
        "item_description": "PVC Rail mats with molded LALO Bottle nubs, 2026 Updated Art",
        "uom": "1/EA",
        "price": "$6.02",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10005_1.png",
        "sort_order": 11
    },
    {
        "item_name": "LALO Bar Mat",
        "brand": "LALO",
        "channel": "On Premise",
        "product_type": "Barware",
        "key_program": "Evergreen",
        "item_description": "PVC Square mats with molded LALO Bottle nubs, 2026 Updated Art",
        "uom": "1/EA",
        "price": "$10.55",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10006_1.png",
        "sort_order": 12
    },
    {
        "item_name": "LALO Bottle Pin",
        "brand": "LALO",
        "channel": "Giveaways",
        "product_type": "Promo Items",
        "key_program": "Sweeps",
        "item_description": "Custom Lalo bottle pins",
        "uom": "1/EA",
        "price": "$0.94",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10007_1.png",
        "sort_order": 13
    },
    {
        "item_name": "LALO TW&M OND Gift Box Display",
        "brand": "LALO",
        "channel": "Off Premise",
        "product_type": "Displays",
        "key_program": "Holiday",
        "item_description": "Holiday gift box display featuring a Lalo Tequila bottle and finished with a metallic ribbon.",
        "uom": "1/EA",
        "price": "$469.60",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_10008_1.png",
        "sort_order": 14
    },
    {
        "item_name": "LALO Shelf Talker - 25/PK",
        "brand": "LALO",
        "channel": "Off Premise",
        "product_type": "Print",
        "key_program": "Evergreen",
        "item_description": "Shelf talker featuring CMYK printing on one side and scored. Includes double sided tape on the back of each unit, positioned at the top center.",
        "uom": "25/PK",
        "price": "$7.17",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_53553_1.jpg?fcts=20260519085204",
        "sort_order": 15
    },
    {
        "item_name": "LALO A Frame Chalkboard Sign",
        "brand": "LALO",
        "channel": "Off Premise",
        "product_type": "Displays",
        "key_program": "Evergreen",
        "item_description": "Lalo bottle silhouette chalkboard made from MDF",
        "uom": "1/EA",
        "price": "$141.80",
        "image_url": "https://demo.mrlsmartbuy.com/images/4006_60045_1.png?fcts=20260519084157",
        "sort_order": 16
    }
]

Role = Literal["admin", "editor"]


class LoginRequest(BaseModel):
    email: EmailStr
    password: str


class ItemCreate(BaseModel):
    sheet: str = "ad-hoc"
    date_or_buy: str = ""
    current_status: str = ""
    visual_reference: str = ""
    brand: str = ""
    program_name: str = ""
    item_name: str = ""
    qty: str = ""
    important_notes: str = ""
    mrl_order_number: str = ""
    estimated_ship_date: str = ""
    estimated_ihd: str = ""
    tracking: str = ""


class ItemUpdate(BaseModel):
    date_or_buy: str | None = None
    current_status: str | None = None
    visual_reference: str | None = None
    brand: str | None = None
    program_name: str | None = None
    item_name: str | None = None
    qty: str | None = None
    important_notes: str | None = None
    mrl_order_number: str | None = None
    estimated_ship_date: str | None = None
    estimated_ihd: str | None = None
    tracking: str | None = None


class ItemBatchUpdate(BaseModel):
    id: int
    changes: ItemUpdate


class ItemChangeBatch(BaseModel):
    creates: list[ItemCreate] = []
    updates: list[ItemBatchUpdate] = []
    deletes: list[int] = []


class UserCreate(BaseModel):
    email: EmailStr
    name: str = ""
    password: str = Field(min_length=6)
    role: Role = "editor"
    is_active: bool = True


class UserUpdate(BaseModel):
    email: EmailStr | None = None
    name: str | None = None
    role: Role | None = None
    is_active: bool | None = None


class PasswordReset(BaseModel):
    password: str = Field(min_length=6)


class KanbanColumnCreate(BaseModel):
    title: str = Field(min_length=1)
    statuses: list[str] = []
    sort_order: int = 0
    is_visible: bool = True


class KanbanColumnUpdate(BaseModel):
    title: str | None = Field(default=None, min_length=1)
    statuses: list[str] | None = None
    sort_order: int | None = None
    is_visible: bool | None = None


class SheetStatusCreate(BaseModel):
    sheet: str
    status: str = Field(min_length=1)


class SurveyResponseCreate(BaseModel):
    survey_item_id: int
    email: EmailStr
    attention_effectiveness: int = Field(ge=1, le=5)
    recommend_rollout: Literal["Yes", "No", "Maybe"]
    retail_engagement: int = Field(ge=1, le=5)
    stands_out: Literal["Yes", "No", "Neutral"]
    price_reasonable: Literal["Yes", "No"]
    feedback: str = ""


def connect() -> Any:
    if IS_POSTGRES:
        if psycopg is None or dict_row is None:
            raise RuntimeError("psycopg is required when DATABASE_URL points to PostgreSQL")
        return psycopg.connect(DATABASE_URL, row_factory=dict_row)

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def db_sql(query: str) -> str:
    if IS_POSTGRES:
        return query.replace("?", "%s")
    return query


def execute(conn: Any, query: str, params: tuple | list = ()) -> Any:
    return conn.execute(db_sql(query), tuple(params))


def ensure_column(conn: Any, table_name: str, column_name: str, column_definition: str) -> None:
    if IS_POSTGRES:
        execute(conn, f"ALTER TABLE {table_name} ADD COLUMN IF NOT EXISTS {column_name} {column_definition}")
        return

    existing_columns = execute(conn, f"PRAGMA table_info({table_name})").fetchall()
    if column_name not in {column["name"] for column in existing_columns}:
        execute(conn, f"ALTER TABLE {table_name} ADD COLUMN {column_name} {column_definition}")


def hash_password(password: str, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120_000)
    return f"{salt}${digest.hex()}"


def verify_password(password: str, password_hash: str) -> bool:
    try:
        salt, _ = password_hash.split("$", 1)
    except ValueError:
        return False
    return hmac.compare_digest(hash_password(password, salt), password_hash)


def utc_now() -> datetime:
    return datetime.now(UTC)


def user_to_public(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "email": row["email"],
        "name": row["name"],
        "role": row["role"],
        "is_active": bool(row["is_active"]),
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def audit_log_to_public(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "user_id": row["user_id"],
        "user_email": row["user_email"],
        "action": row["action"],
        "sheet": row["sheet"],
        "item_id": row["item_id"],
        "before": json.loads(row["before_json"] or "null"),
        "after": json.loads(row["after_json"] or "null"),
        "created_at": row["created_at"],
    }


def kanban_column_to_public(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "title": row["title"],
        "statuses": json.loads(row["statuses_json"] or "[]"),
        "sort_order": row["sort_order"],
        "is_visible": bool(row["is_visible"]),
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def sheet_status_to_public(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "sheet": row["sheet"],
        "status": row["status"],
        "sort_order": row["sort_order"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def survey_item_to_public(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "item_name": row["item_name"],
        "brand": row["brand"],
        "channel": row["channel"],
        "product_type": row["product_type"],
        "key_program": row["key_program"],
        "item_description": row["item_description"],
        "uom": row["uom"],
        "price": row["price"],
        "image_url": row["image_url"],
        "sort_order": row["sort_order"],
        "is_active": bool(row["is_active"]),
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def survey_response_to_public(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "survey_item_id": row["survey_item_id"],
        "email": row["email"],
        "attention_effectiveness": row["attention_effectiveness"],
        "recommend_rollout": row["recommend_rollout"],
        "retail_engagement": row["retail_engagement"],
        "stands_out": row["stands_out"],
        "price_reasonable": row["price_reasonable"],
        "feedback": row["feedback"],
        "created_at": row["created_at"],
    }


def write_audit_log(
    conn: Any,
    user: sqlite3.Row,
    action: str,
    sheet: str | None = None,
    item_id: int | None = None,
    before: dict | None = None,
    after: dict | None = None,
) -> None:
    execute(
        conn,
        """
        INSERT INTO audit_logs (
            user_id, user_email, action, sheet, item_id, before_json, after_json
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            user["id"],
            user["email"],
            action,
            sheet,
            item_id,
            json.dumps(before, default=str) if before is not None else None,
            json.dumps(after, default=str) if after is not None else None,
        ),
    )


def seed_tracker_items_if_empty(conn: Any) -> None:
    if not SEED_ITEMS_PATH.exists():
        return

    row_count = execute(conn, "SELECT COUNT(*) AS row_count FROM tracker_items").fetchone()
    if row_count["row_count"] > 0:
        return

    seed_items = json.loads(SEED_ITEMS_PATH.read_text(encoding="utf-8"))
    for item in seed_items:
        execute(
            conn,
            """
            INSERT INTO tracker_items (
                sheet, source_row, sort_order, date_or_buy, current_status,
                visual_reference, brand, program_name, item_name, qty,
                important_notes, mrl_order_number, estimated_ship_date,
                estimated_ihd, tracking, extra_json
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                item.get("sheet", "ad-hoc"),
                item.get("source_row"),
                item.get("sort_order", 0),
                item.get("date_or_buy", ""),
                item.get("current_status", ""),
                item.get("visual_reference", ""),
                item.get("brand", ""),
                item.get("program_name", ""),
                item.get("item_name", ""),
                item.get("qty", ""),
                item.get("important_notes", ""),
                item.get("mrl_order_number", ""),
                item.get("estimated_ship_date", ""),
                item.get("estimated_ihd", ""),
                item.get("tracking", ""),
                item.get("extra_json", "{}"),
            ),
        )


def sync_seeded_tracker_item_catalog(conn: Any) -> None:
    if not SEED_ITEMS_PATH.exists():
        return

    seed_items = json.loads(SEED_ITEMS_PATH.read_text(encoding="utf-8"))
    for item in seed_items:
        source_row = item.get("source_row")
        if source_row is None:
            continue

        execute(
            conn,
            """
            UPDATE tracker_items
            SET visual_reference = ?, brand = ?, program_name = ?, item_name = ?, extra_json = ?
            WHERE sheet = ? AND source_row = ?
            """,
            (
                item.get("visual_reference", ""),
                item.get("brand", ""),
                item.get("program_name", ""),
                item.get("item_name", ""),
                item.get("extra_json", "{}"),
                item.get("sheet", "ad-hoc"),
                source_row,
            ),
        )


def seed_kanban_columns_if_empty(conn: Any) -> None:
    row_count = execute(conn, "SELECT COUNT(*) AS row_count FROM kanban_columns").fetchone()
    if row_count["row_count"] > 0:
        return

    for column in DEFAULT_KANBAN_COLUMNS:
        execute(
            conn,
            """
            INSERT INTO kanban_columns (title, statuses_json, sort_order, is_visible)
            VALUES (?, ?, ?, ?)
            """,
            (
                column["title"],
                json.dumps(column["statuses"]),
                column["sort_order"],
                True,
            ),
        )


def seed_sheet_statuses_if_empty(conn: Any) -> None:
    row_count = execute(conn, "SELECT COUNT(*) AS row_count FROM sheet_statuses").fetchone()
    if row_count["row_count"] > 0:
        return

    for sheet_key, sheet_config in SHEETS.items():
        for sort_order, status_label in enumerate(sheet_config["statuses"], start=1):
            execute(
                conn,
                """
                INSERT INTO sheet_statuses (sheet, status, sort_order)
                VALUES (?, ?, ?)
                """,
                (sheet_key, status_label, sort_order),
            )


def sync_default_survey_items(conn: Any) -> None:
    active_item_names = []

    for item in DEFAULT_SURVEY_ITEMS:
        active_item_names.append(item["item_name"].lower())
        existing_item = execute(
            conn,
            "SELECT id FROM survey_items WHERE lower(item_name) = lower(?) ORDER BY id ASC LIMIT 1",
            (item["item_name"],),
        ).fetchone()

        values = (
            item["item_name"],
            item["brand"],
            item["channel"],
            item["product_type"],
            item["key_program"],
            item["item_description"],
            item["uom"],
            item["price"],
            item["image_url"],
            item["sort_order"],
            True,
        )
        if existing_item is None:
            execute(
                conn,
                """
                INSERT INTO survey_items (
                    item_name, brand, channel, product_type, key_program,
                    item_description, uom, price, image_url, sort_order, is_active
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
        else:
            execute(
                conn,
                """
                UPDATE survey_items
                SET item_name = ?, brand = ?, channel = ?, product_type = ?, key_program = ?,
                    item_description = ?, uom = ?, price = ?, image_url = ?, sort_order = ?, is_active = ?
                WHERE id = ?
                """,
                (*values, existing_item["id"]),
            )

    placeholders = ", ".join("?" for _ in active_item_names)
    execute(
        conn,
        f"UPDATE survey_items SET is_active = ? WHERE lower(item_name) NOT IN ({placeholders})",
        (False, *active_item_names),
    )


def get_sheet_config() -> dict:
    sheets = {
        sheet_key: {
            "label": sheet_config["label"],
            "dateLabel": sheet_config["dateLabel"],
            "statuses": [],
        }
        for sheet_key, sheet_config in SHEETS.items()
    }
    with connect() as conn:
        rows = execute(
            conn,
            "SELECT * FROM sheet_statuses ORDER BY sheet ASC, sort_order ASC, id ASC",
        ).fetchall()

    for row in rows:
        if row["sheet"] in sheets:
            sheets[row["sheet"]]["statuses"].append(row["status"])

    return sheets


def init_db() -> None:
    with connect() as conn:
        if IS_POSTGRES:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS tracker_items (
                    id SERIAL PRIMARY KEY,
                    sheet TEXT NOT NULL,
                    source_row INTEGER,
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    date_or_buy TEXT,
                    current_status TEXT,
                    visual_reference TEXT,
                    brand TEXT,
                    program_name TEXT,
                    item_name TEXT,
                    qty TEXT,
                    important_notes TEXT,
                    mrl_order_number TEXT,
                    estimated_ship_date TEXT,
                    estimated_ihd TEXT,
                    tracking TEXT,
                    extra_json TEXT NOT NULL DEFAULT '{}',
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id SERIAL PRIMARY KEY,
                    email TEXT NOT NULL UNIQUE,
                    name TEXT NOT NULL DEFAULT '',
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL CHECK(role IN ('admin', 'editor')),
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS sessions (
                    token TEXT PRIMARY KEY,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    expires_at TEXT NOT NULL,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS audit_logs (
                    id SERIAL PRIMARY KEY,
                    user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    user_email TEXT NOT NULL,
                    action TEXT NOT NULL,
                    sheet TEXT,
                    item_id INTEGER,
                    before_json TEXT,
                    after_json TEXT,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS kanban_columns (
                    id SERIAL PRIMARY KEY,
                    title TEXT NOT NULL,
                    statuses_json TEXT NOT NULL DEFAULT '[]',
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    is_visible BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS sheet_statuses (
                    id SERIAL PRIMARY KEY,
                    sheet TEXT NOT NULL,
                    status TEXT NOT NULL,
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(sheet, status)
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS survey_items (
                    id SERIAL PRIMARY KEY,
                    item_name TEXT NOT NULL,
                    brand TEXT NOT NULL DEFAULT '',
                    channel TEXT NOT NULL DEFAULT '',
                    product_type TEXT NOT NULL DEFAULT '',
                    key_program TEXT NOT NULL DEFAULT '',
                    item_description TEXT NOT NULL DEFAULT '',
                    uom TEXT NOT NULL DEFAULT '',
                    price TEXT NOT NULL DEFAULT '',
                    image_url TEXT NOT NULL DEFAULT '',
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS survey_responses (
                    id SERIAL PRIMARY KEY,
                    survey_item_id INTEGER NOT NULL REFERENCES survey_items(id) ON DELETE CASCADE,
                    email TEXT NOT NULL,
                    attention_effectiveness INTEGER NOT NULL,
                    recommend_rollout TEXT NOT NULL,
                    retail_engagement INTEGER NOT NULL,
                    stands_out TEXT NOT NULL,
                    price_reasonable TEXT NOT NULL,
                    feedback TEXT NOT NULL DEFAULT '',
                    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
        else:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS tracker_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    sheet TEXT NOT NULL,
                    source_row INTEGER,
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    date_or_buy TEXT,
                    current_status TEXT,
                    visual_reference TEXT,
                    brand TEXT,
                    program_name TEXT,
                    item_name TEXT,
                    qty TEXT,
                    important_notes TEXT,
                    mrl_order_number TEXT,
                    estimated_ship_date TEXT,
                    estimated_ihd TEXT,
                    tracking TEXT,
                    extra_json TEXT NOT NULL DEFAULT '{}',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    email TEXT NOT NULL UNIQUE,
                    name TEXT NOT NULL DEFAULT '',
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL CHECK(role IN ('admin', 'editor')),
                    is_active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS sessions (
                    token TEXT PRIMARY KEY,
                    user_id INTEGER NOT NULL,
                    expires_at TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS audit_logs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER,
                    user_email TEXT NOT NULL,
                    action TEXT NOT NULL,
                    sheet TEXT,
                    item_id INTEGER,
                    before_json TEXT,
                    after_json TEXT,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE SET NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS kanban_columns (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    title TEXT NOT NULL,
                    statuses_json TEXT NOT NULL DEFAULT '[]',
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    is_visible INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS sheet_statuses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    sheet TEXT NOT NULL,
                    status TEXT NOT NULL,
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(sheet, status)
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS survey_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    item_name TEXT NOT NULL,
                    brand TEXT NOT NULL DEFAULT '',
                    channel TEXT NOT NULL DEFAULT '',
                    product_type TEXT NOT NULL DEFAULT '',
                    key_program TEXT NOT NULL DEFAULT '',
                    item_description TEXT NOT NULL DEFAULT '',
                    uom TEXT NOT NULL DEFAULT '',
                    price TEXT NOT NULL DEFAULT '',
                    image_url TEXT NOT NULL DEFAULT '',
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    is_active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS survey_responses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    survey_item_id INTEGER NOT NULL,
                    email TEXT NOT NULL,
                    attention_effectiveness INTEGER NOT NULL,
                    recommend_rollout TEXT NOT NULL,
                    retail_engagement INTEGER NOT NULL,
                    stands_out TEXT NOT NULL,
                    price_reasonable TEXT NOT NULL,
                    feedback TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (survey_item_id) REFERENCES survey_items(id) ON DELETE CASCADE
                )
                """
            )

        execute(conn, "DELETE FROM sessions WHERE expires_at <= ?", (utc_now().isoformat(),))
        ensure_column(conn, "survey_items", "product_type", "TEXT NOT NULL DEFAULT ''")
        ensure_column(conn, "survey_items", "key_program", "TEXT NOT NULL DEFAULT ''")
        existing_admin = execute(
            conn,
            "SELECT id FROM users WHERE lower(email) = lower(?)",
            (ADMIN_EMAIL,),
        ).fetchone()
        if existing_admin is None:
            execute(
                conn,
                """
                INSERT INTO users (email, name, password_hash, role)
                VALUES (?, ?, ?, 'admin')
                """,
                (ADMIN_EMAIL, ADMIN_NAME, hash_password(ADMIN_PASSWORD)),
            )
        seed_tracker_items_if_empty(conn)
        sync_seeded_tracker_item_catalog(conn)
        seed_kanban_columns_if_empty(conn)
        seed_sheet_statuses_if_empty(conn)
        sync_default_survey_items(conn)


def row_to_item(row: sqlite3.Row) -> dict:
    item = dict(row)
    item["extra"] = json.loads(item.pop("extra_json") or "{}")
    return item


init_db()
app = FastAPI(title="Phillips Project Tracker API", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_current_user(authorization: str | None = Header(default=None)) -> sqlite3.Row:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Missing login token")

    token = authorization.split(" ", 1)[1].strip()
    with connect() as conn:
        row = execute(
            conn,
            """
            SELECT users.*
            FROM sessions
            JOIN users ON users.id = sessions.user_id
            WHERE sessions.token = ? AND sessions.expires_at > ?
            """,
            (token, utc_now().isoformat()),
        ).fetchone()

    if row is None or not row["is_active"]:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid or expired login token")
    return row


def require_admin(user: sqlite3.Row = Depends(get_current_user)) -> sqlite3.Row:
    if user["role"] != "admin":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Admin access required")
    return user


@app.get("/api/health")
def health() -> dict:
    return {"ok": True}


@app.post("/api/auth/login")
def login(payload: LoginRequest) -> dict:
    with connect() as conn:
        user = execute(
            conn,
            "SELECT * FROM users WHERE lower(email) = lower(?)",
            (payload.email,),
        ).fetchone()
        if user is None or not verify_password(payload.password, user["password_hash"]):
            raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid email or password")
        if not user["is_active"]:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Account is inactive")

        token = secrets.token_urlsafe(32)
        expires_at = (utc_now() + timedelta(days=SESSION_DAYS)).isoformat()
        execute(
            conn,
            "INSERT INTO sessions (token, user_id, expires_at) VALUES (?, ?, ?)",
            (token, user["id"], expires_at),
        )

    return {"token": token, "user": user_to_public(user)}


@app.post("/api/auth/logout", status_code=status.HTTP_204_NO_CONTENT)
def logout(
    response: Response,
    authorization: str | None = Header(default=None),
    user: sqlite3.Row = Depends(get_current_user),
) -> Response:
    del user
    if authorization:
        token = authorization.split(" ", 1)[1].strip()
        with connect() as conn:
            execute(conn, "DELETE FROM sessions WHERE token = ?", (token,))
    response.status_code = status.HTTP_204_NO_CONTENT
    return response


@app.get("/api/auth/me")
def me(user: sqlite3.Row = Depends(get_current_user)) -> dict:
    return {"user": user_to_public(user)}


@app.get("/api/sheets")
def get_sheets(user: sqlite3.Row = Depends(get_current_user)) -> dict:
    del user
    return {"sheets": get_sheet_config()}


@app.get("/api/public/items")
def list_public_items(q: str = "") -> dict:
    clauses = ["sheet IN (?, ?)", "COALESCE(current_status, '') != ?"]
    args: list[str] = ["ad-hoc", "buys", "Delivered"]

    if q.strip():
        like = f"%{q.strip().lower()}%"
        clauses.append(
            "("
            "lower(brand) LIKE ? OR lower(program_name) LIKE ? OR "
            "lower(item_name) LIKE ? OR lower(important_notes) LIKE ? OR "
            "lower(mrl_order_number) LIKE ? OR lower(tracking) LIKE ? OR "
            "lower(extra_json) LIKE ?"
            ")"
        )
        args.extend([like] * 7)

    with connect() as conn:
        rows = execute(
            conn,
            f"""
            SELECT * FROM tracker_items
            WHERE {' AND '.join(clauses)}
            ORDER BY sheet ASC, sort_order ASC, id ASC
            """,
            args,
        ).fetchall()
    return {"items": [row_to_item(row) for row in rows]}


@app.get("/api/public/kanban-columns")
def list_public_kanban_columns() -> dict:
    with connect() as conn:
        rows = execute(
            conn,
            """
            SELECT * FROM kanban_columns
            WHERE is_visible = ?
            ORDER BY sort_order ASC, id ASC
            """,
            (True,),
        ).fetchall()
    return {"columns": [kanban_column_to_public(row) for row in rows]}


@app.get("/api/public/survey-items")
def list_public_survey_items(q: str = "") -> dict:
    clauses = ["is_active = ?"]
    args: list[Any] = [True]
    if q.strip():
        like = f"%{q.strip().lower()}%"
        clauses.append(
            "("
            "lower(item_name) LIKE ? OR lower(brand) LIKE ? OR lower(channel) LIKE ? OR "
            "lower(product_type) LIKE ? OR lower(key_program) LIKE ? OR lower(item_description) LIKE ?"
            ")"
        )
        args.extend([like] * 6)

    with connect() as conn:
        rows = execute(
            conn,
            f"""
            SELECT * FROM survey_items
            WHERE {' AND '.join(clauses)}
            ORDER BY sort_order ASC, id ASC
            """,
            args,
        ).fetchall()
    return {"items": [survey_item_to_public(row) for row in rows]}


@app.post("/api/public/survey-responses", status_code=status.HTTP_201_CREATED)
def create_public_survey_response(payload: SurveyResponseCreate) -> dict:
    with connect() as conn:
        survey_item = execute(
            conn,
            "SELECT id FROM survey_items WHERE id = ? AND is_active = ?",
            (payload.survey_item_id, True),
        ).fetchone()
        if survey_item is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Survey item not found")

        values = (
            payload.survey_item_id,
            str(payload.email).lower(),
            payload.attention_effectiveness,
            payload.recommend_rollout,
            payload.retail_engagement,
            payload.stands_out,
            payload.price_reasonable,
            payload.feedback.strip(),
        )
        if IS_POSTGRES:
            row = execute(
                conn,
                """
                INSERT INTO survey_responses (
                    survey_item_id, email, attention_effectiveness, recommend_rollout,
                    retail_engagement, stands_out, price_reasonable, feedback
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                RETURNING *
                """,
                values,
            ).fetchone()
        else:
            cursor = execute(
                conn,
                """
                INSERT INTO survey_responses (
                    survey_item_id, email, attention_effectiveness, recommend_rollout,
                    retail_engagement, stands_out, price_reasonable, feedback
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
            row = execute(conn, "SELECT * FROM survey_responses WHERE id = ?", (cursor.lastrowid,)).fetchone()
    return {"response": survey_response_to_public(row)}


@app.post("/api/uploads/visual-reference")
async def upload_visual_reference(
    file: UploadFile = File(...),
    user: sqlite3.Row = Depends(get_current_user),
) -> dict:
    del user
    if cloudinary is None:
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "Cloudinary dependency is not installed")
    if not (CLOUDINARY_CLOUD_NAME and CLOUDINARY_API_KEY and CLOUDINARY_API_SECRET):
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "Cloudinary is not configured")
    if file.content_type not in ALLOWED_UPLOAD_TYPES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Upload a JPG, PNG, or WEBP image")

    contents = await file.read()
    if len(contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Image must be 8 MB or smaller")

    try:
        result = cloudinary.uploader.upload(
            io.BytesIO(contents),
            folder="phillips-tracker/visual-references",
            resource_type="image",
        )
    except Exception as exc:
        raise HTTPException(
            status.HTTP_502_BAD_GATEWAY,
            "Cloudinary upload failed. Check the Cloudinary cloud name, API key, and API secret.",
        ) from exc
    return {
        "url": result["secure_url"],
        "public_id": result["public_id"],
    }


@app.get("/api/items")
def list_items(
    sheet: str = "ad-hoc",
    q: str = "",
    status_filter: str = Query("", alias="status"),
    user: sqlite3.Row = Depends(get_current_user),
) -> dict:
    del user
    if sheet not in SHEETS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown sheet")

    clauses = ["sheet = ?"]
    args: list[str] = [sheet]
    if status_filter:
        clauses.append("current_status = ?")
        args.append(status_filter)
    if q.strip():
        like = f"%{q.strip().lower()}%"
        clauses.append(
            "("
            "lower(brand) LIKE ? OR lower(program_name) LIKE ? OR "
            "lower(item_name) LIKE ? OR lower(important_notes) LIKE ? OR "
            "lower(mrl_order_number) LIKE ? OR lower(tracking) LIKE ?"
            ")"
        )
        args.extend([like] * 6)

    with connect() as conn:
        rows = execute(
            conn,
            f"""
            SELECT * FROM tracker_items
            WHERE {' AND '.join(clauses)}
            ORDER BY sort_order ASC, id ASC
            """,
            args,
        ).fetchall()
    return {"items": [row_to_item(row) for row in rows]}


@app.post("/api/items", status_code=status.HTTP_201_CREATED)
def create_item(payload: ItemCreate, user: sqlite3.Row = Depends(get_current_user)) -> dict:
    if payload.sheet not in SHEETS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown sheet")

    with connect() as conn:
        max_sort_row = execute(
            conn,
            "SELECT COALESCE(MAX(sort_order), 0) AS max_sort FROM tracker_items WHERE sheet = ?",
            (payload.sheet,),
        ).fetchone()
        max_sort = max_sort_row["max_sort"]
        values = (
            payload.sheet,
            max_sort + 1,
            payload.date_or_buy,
            payload.current_status,
            payload.visual_reference,
            payload.brand,
            payload.program_name,
            payload.item_name,
            payload.qty,
            payload.important_notes,
            payload.mrl_order_number,
            payload.estimated_ship_date,
            payload.estimated_ihd,
            payload.tracking,
        )
        if IS_POSTGRES:
            row = execute(
                conn,
                """
                INSERT INTO tracker_items (
                    sheet, sort_order, date_or_buy, current_status, visual_reference,
                    brand, program_name, item_name, qty, important_notes,
                    mrl_order_number, estimated_ship_date, estimated_ihd, tracking
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                RETURNING *
                """,
                values,
            ).fetchone()
        else:
            cursor = execute(
                conn,
                """
                INSERT INTO tracker_items (
                    sheet, sort_order, date_or_buy, current_status, visual_reference,
                    brand, program_name, item_name, qty, important_notes,
                    mrl_order_number, estimated_ship_date, estimated_ihd, tracking
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
            row = execute(conn, "SELECT * FROM tracker_items WHERE id = ?", (cursor.lastrowid,)).fetchone()
        write_audit_log(
            conn,
            user,
            "create",
            sheet=payload.sheet,
            item_id=row["id"],
            after=row_to_item(row),
        )
    return {"item": row_to_item(row)}


@app.patch("/api/items/{item_id}")
def update_item(item_id: int, payload: ItemUpdate, user: sqlite3.Row = Depends(get_current_user)) -> dict:
    updates = payload.model_dump(exclude_unset=True)
    updates = {key: "" if value is None else str(value) for key, value in updates.items()}
    if not updates:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No editable fields supplied")

    assignments = ", ".join(f"{field} = ?" for field in updates)
    values = list(updates.values())
    values.append(item_id)

    with connect() as conn:
        before_row = execute(conn, "SELECT * FROM tracker_items WHERE id = ?", (item_id,)).fetchone()
        if before_row is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found")
        cursor = execute(
            conn,
            f"""
            UPDATE tracker_items
            SET {assignments}, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            values,
        )
        if cursor.rowcount == 0:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found")
        row = execute(conn, "SELECT * FROM tracker_items WHERE id = ?", (item_id,)).fetchone()
        write_audit_log(
            conn,
            user,
            "update",
            sheet=row["sheet"],
            item_id=item_id,
            before=row_to_item(before_row),
            after=row_to_item(row),
        )
    return {"item": row_to_item(row)}


@app.delete("/api/items/{item_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_item(
    item_id: int,
    response: Response,
    user: sqlite3.Row = Depends(require_admin),
) -> Response:
    with connect() as conn:
        before_row = execute(conn, "SELECT * FROM tracker_items WHERE id = ?", (item_id,)).fetchone()
        if before_row is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found")
        cursor = execute(conn, "DELETE FROM tracker_items WHERE id = ?", (item_id,))
        write_audit_log(
            conn,
            user,
            "delete",
            sheet=before_row["sheet"],
            item_id=item_id,
            before=row_to_item(before_row),
        )
    if cursor.rowcount == 0:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found")
    response.status_code = status.HTTP_204_NO_CONTENT
    return response


@app.post("/api/items/changes")
def submit_item_changes(
    payload: ItemChangeBatch,
    user: sqlite3.Row = Depends(get_current_user),
) -> dict:
    created = 0
    updated = 0
    deleted = 0

    with connect() as conn:
        for create_payload in payload.creates:
            if create_payload.sheet not in SHEETS:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown sheet")
            max_sort_row = execute(
                conn,
                "SELECT COALESCE(MAX(sort_order), 0) AS max_sort FROM tracker_items WHERE sheet = ?",
                (create_payload.sheet,),
            ).fetchone()
            values = (
                create_payload.sheet,
                max_sort_row["max_sort"] + 1,
                create_payload.date_or_buy,
                create_payload.current_status,
                create_payload.visual_reference,
                create_payload.brand,
                create_payload.program_name,
                create_payload.item_name,
                create_payload.qty,
                create_payload.important_notes,
                create_payload.mrl_order_number,
                create_payload.estimated_ship_date,
                create_payload.estimated_ihd,
                create_payload.tracking,
            )
            if IS_POSTGRES:
                row = execute(
                    conn,
                    """
                    INSERT INTO tracker_items (
                        sheet, sort_order, date_or_buy, current_status, visual_reference,
                        brand, program_name, item_name, qty, important_notes,
                        mrl_order_number, estimated_ship_date, estimated_ihd, tracking
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    RETURNING *
                    """,
                    values,
                ).fetchone()
            else:
                cursor = execute(
                    conn,
                    """
                    INSERT INTO tracker_items (
                        sheet, sort_order, date_or_buy, current_status, visual_reference,
                        brand, program_name, item_name, qty, important_notes,
                        mrl_order_number, estimated_ship_date, estimated_ihd, tracking
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    values,
                )
                row = execute(conn, "SELECT * FROM tracker_items WHERE id = ?", (cursor.lastrowid,)).fetchone()
            created += 1
            write_audit_log(
                conn,
                user,
                "create",
                sheet=create_payload.sheet,
                item_id=row["id"],
                after=row_to_item(row),
            )

        for update_payload in payload.updates:
            updates = update_payload.changes.model_dump(exclude_unset=True)
            updates = {key: "" if value is None else str(value) for key, value in updates.items()}
            if not updates:
                continue
            before_row = execute(
                conn,
                "SELECT * FROM tracker_items WHERE id = ?",
                (update_payload.id,),
            ).fetchone()
            if before_row is None:
                raise HTTPException(status.HTTP_404_NOT_FOUND, "Item not found")
            assignments = ", ".join(f"{field} = ?" for field in updates)
            values = list(updates.values())
            values.append(update_payload.id)
            execute(
                conn,
                f"""
                UPDATE tracker_items
                SET {assignments}, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                values,
            )
            after_row = execute(
                conn,
                "SELECT * FROM tracker_items WHERE id = ?",
                (update_payload.id,),
            ).fetchone()
            updated += 1
            write_audit_log(
                conn,
                user,
                "update",
                sheet=after_row["sheet"],
                item_id=update_payload.id,
                before=row_to_item(before_row),
                after=row_to_item(after_row),
            )

        for item_id in payload.deletes:
            before_row = execute(conn, "SELECT * FROM tracker_items WHERE id = ?", (item_id,)).fetchone()
            if before_row is None:
                continue
            execute(conn, "DELETE FROM tracker_items WHERE id = ?", (item_id,))
            deleted += 1
            write_audit_log(
                conn,
                user,
                "delete",
                sheet=before_row["sheet"],
                item_id=item_id,
                before=row_to_item(before_row),
            )

    return {"ok": True, "created": created, "updated": updated, "deleted": deleted}


@app.get("/api/admin/users")
def list_users(user: sqlite3.Row = Depends(require_admin)) -> dict:
    del user
    with connect() as conn:
        rows = execute(conn, "SELECT * FROM users ORDER BY created_at DESC, id DESC").fetchall()
    return {"users": [user_to_public(row) for row in rows]}


@app.get("/api/admin/kanban-columns")
def list_admin_kanban_columns(user: sqlite3.Row = Depends(require_admin)) -> dict:
    del user
    with connect() as conn:
        rows = execute(
            conn,
            "SELECT * FROM kanban_columns ORDER BY sort_order ASC, id ASC",
        ).fetchall()
    return {"columns": [kanban_column_to_public(row) for row in rows]}


@app.post("/api/admin/kanban-columns", status_code=status.HTTP_201_CREATED)
def create_kanban_column(
    payload: KanbanColumnCreate,
    user: sqlite3.Row = Depends(require_admin),
) -> dict:
    del user
    with connect() as conn:
        values = (
            payload.title.strip(),
            json.dumps(payload.statuses),
            payload.sort_order,
            payload.is_visible,
        )
        if IS_POSTGRES:
            row = execute(
                conn,
                """
                INSERT INTO kanban_columns (title, statuses_json, sort_order, is_visible)
                VALUES (?, ?, ?, ?)
                RETURNING *
                """,
                values,
            ).fetchone()
        else:
            cursor = execute(
                conn,
                """
                INSERT INTO kanban_columns (title, statuses_json, sort_order, is_visible)
                VALUES (?, ?, ?, ?)
                """,
                values,
            )
            row = execute(conn, "SELECT * FROM kanban_columns WHERE id = ?", (cursor.lastrowid,)).fetchone()
    return {"column": kanban_column_to_public(row)}


@app.patch("/api/admin/kanban-columns/{column_id}")
def update_kanban_column(
    column_id: int,
    payload: KanbanColumnUpdate,
    user: sqlite3.Row = Depends(require_admin),
) -> dict:
    del user
    updates = payload.model_dump(exclude_unset=True)
    if "title" in updates and updates["title"] is not None:
        updates["title"] = updates["title"].strip()
    if "statuses" in updates:
        updates["statuses_json"] = json.dumps(updates.pop("statuses") or [])
    if not updates:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No kanban column fields supplied")

    assignments = ", ".join(f"{field} = ?" for field in updates)
    values = list(updates.values())
    values.append(column_id)
    with connect() as conn:
        cursor = execute(
            conn,
            f"""
            UPDATE kanban_columns
            SET {assignments}, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            values,
        )
        if cursor.rowcount == 0:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Kanban column not found")
        row = execute(conn, "SELECT * FROM kanban_columns WHERE id = ?", (column_id,)).fetchone()
    return {"column": kanban_column_to_public(row)}


@app.delete("/api/admin/kanban-columns/{column_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_kanban_column(
    column_id: int,
    response: Response,
    user: sqlite3.Row = Depends(require_admin),
) -> Response:
    del user
    with connect() as conn:
        cursor = execute(conn, "DELETE FROM kanban_columns WHERE id = ?", (column_id,))
    if cursor.rowcount == 0:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Kanban column not found")
    response.status_code = status.HTTP_204_NO_CONTENT
    return response


@app.get("/api/admin/sheet-statuses")
def list_sheet_statuses(user: sqlite3.Row = Depends(require_admin)) -> dict:
    del user
    with connect() as conn:
        rows = execute(
            conn,
            "SELECT * FROM sheet_statuses ORDER BY sheet ASC, sort_order ASC, id ASC",
        ).fetchall()
    return {"statuses": [sheet_status_to_public(row) for row in rows]}


@app.post("/api/admin/sheet-statuses", status_code=status.HTTP_201_CREATED)
def create_sheet_status(
    payload: SheetStatusCreate,
    user: sqlite3.Row = Depends(require_admin),
) -> dict:
    del user
    if payload.sheet not in SHEETS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unknown sheet")

    status_label = payload.status.strip()
    if not status_label:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Status is required")

    try:
        with connect() as conn:
            max_sort_row = execute(
                conn,
                "SELECT COALESCE(MAX(sort_order), 0) AS max_sort FROM sheet_statuses WHERE sheet = ?",
                (payload.sheet,),
            ).fetchone()
            values = (payload.sheet, status_label, max_sort_row["max_sort"] + 1)
            if IS_POSTGRES:
                row = execute(
                    conn,
                    """
                    INSERT INTO sheet_statuses (sheet, status, sort_order)
                    VALUES (?, ?, ?)
                    RETURNING *
                    """,
                    values,
                ).fetchone()
            else:
                cursor = execute(
                    conn,
                    """
                    INSERT INTO sheet_statuses (sheet, status, sort_order)
                    VALUES (?, ?, ?)
                    """,
                    values,
                )
                row = execute(conn, "SELECT * FROM sheet_statuses WHERE id = ?", (cursor.lastrowid,)).fetchone()
    except DB_INTEGRITY_ERRORS:
        raise HTTPException(status.HTTP_409_CONFLICT, "That status already exists for this sheet")
    return {"status": sheet_status_to_public(row)}


@app.delete("/api/admin/sheet-statuses/{status_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_sheet_status(
    status_id: int,
    response: Response,
    user: sqlite3.Row = Depends(require_admin),
) -> Response:
    del user
    with connect() as conn:
        cursor = execute(conn, "DELETE FROM sheet_statuses WHERE id = ?", (status_id,))
    if cursor.rowcount == 0:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Status not found")
    response.status_code = status.HTTP_204_NO_CONTENT
    return response


@app.get("/api/admin/audit-logs")
def list_audit_logs(user: sqlite3.Row = Depends(require_admin)) -> dict:
    del user
    with connect() as conn:
        rows = execute(
            conn,
            """
            SELECT * FROM audit_logs
            ORDER BY created_at DESC, id DESC
            LIMIT 250
            """,
        ).fetchall()
    return {"logs": [audit_log_to_public(row) for row in rows]}


@app.get("/api/admin/survey-responses/export")
def export_survey_responses(user: sqlite3.Row = Depends(require_admin)) -> Response:
    del user
    with connect() as conn:
        rows = execute(
            conn,
            """
            SELECT
                survey_responses.created_at,
                COALESCE(survey_items.item_name, '') AS item_name,
                COALESCE(survey_items.brand, '') AS brand,
                COALESCE(survey_items.channel, '') AS channel,
                COALESCE(survey_items.product_type, '') AS product_type,
                COALESCE(survey_items.key_program, '') AS key_program,
                COALESCE(survey_items.uom, '') AS uom,
                COALESCE(survey_items.price, '') AS price,
                survey_responses.email,
                survey_responses.recommend_rollout,
                survey_responses.feedback
            FROM survey_responses
            LEFT JOIN survey_items ON survey_items.id = survey_responses.survey_item_id
            ORDER BY survey_responses.created_at DESC, survey_responses.id DESC
            """,
        ).fetchall()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Survey Responses"
    headers = [
        "Submitted At",
        "Item Name",
        "Brand",
        "Channel",
        "Product Type",
        "Key Program",
        "UOM",
        "Price",
        "Email",
        "Interest",
        "Notes",
    ]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        worksheet.append(
            [
                str(row["created_at"] or ""),
                row["item_name"],
                row["brand"],
                row["channel"],
                row["product_type"],
                row["key_program"],
                row["uom"],
                row["price"],
                row["email"],
                "Interested" if row["recommend_rollout"] == "Yes" else "Not Interested",
                row["feedback"],
            ]
        )

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 55)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"survey-responses-{utc_now().date().isoformat()}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/admin/users", status_code=status.HTTP_201_CREATED)
def create_user(payload: UserCreate, user: sqlite3.Row = Depends(require_admin)) -> dict:
    del user
    try:
        with connect() as conn:
            values = (
                str(payload.email).lower(),
                payload.name,
                hash_password(payload.password),
                payload.role,
                payload.is_active,
            )
            if IS_POSTGRES:
                row = execute(
                    conn,
                    """
                    INSERT INTO users (email, name, password_hash, role, is_active)
                    VALUES (?, ?, ?, ?, ?)
                    RETURNING *
                    """,
                    values,
                ).fetchone()
            else:
                cursor = execute(
                    conn,
                    """
                    INSERT INTO users (email, name, password_hash, role, is_active)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    values,
                )
                row = execute(conn, "SELECT * FROM users WHERE id = ?", (cursor.lastrowid,)).fetchone()
    except DB_INTEGRITY_ERRORS:
        raise HTTPException(status.HTTP_409_CONFLICT, "A user with that email already exists")
    return {"user": user_to_public(row)}


@app.patch("/api/admin/users/{user_id}")
def update_user(
    user_id: int,
    payload: UserUpdate,
    user: sqlite3.Row = Depends(require_admin),
) -> dict:
    if user_id == user["id"] and payload.is_active is False:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "You cannot deactivate your own account")

    updates = payload.model_dump(exclude_unset=True)
    if "email" in updates and updates["email"] is not None:
        updates["email"] = str(updates["email"]).lower()
    if not updates:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No user fields supplied")

    assignments = ", ".join(f"{field} = ?" for field in updates)
    values = list(updates.values())
    values.append(user_id)

    try:
        with connect() as conn:
            cursor = execute(
                conn,
                f"""
                UPDATE users
                SET {assignments}, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                values,
            )
            if cursor.rowcount == 0:
                raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
            row = execute(conn, "SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    except DB_INTEGRITY_ERRORS:
        raise HTTPException(status.HTTP_409_CONFLICT, "A user with that email already exists")
    return {"user": user_to_public(row)}


@app.post("/api/admin/users/{user_id}/reset-password")
def reset_password(
    user_id: int,
    payload: PasswordReset,
    user: sqlite3.Row = Depends(require_admin),
) -> dict:
    del user
    with connect() as conn:
        cursor = execute(
            conn,
            """
            UPDATE users
            SET password_hash = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (hash_password(payload.password), user_id),
        )
        if cursor.rowcount == 0:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "User not found")
        execute(conn, "DELETE FROM sessions WHERE user_id = ?", (user_id,))
    return {"ok": True}


if FRONTEND_DIST.exists():
    assets_dir = FRONTEND_DIST / "assets"
    if assets_dir.exists():
        app.mount("/assets", StaticFiles(directory=assets_dir), name="assets")

    @app.get("/{full_path:path}", include_in_schema=False)
    def serve_frontend(full_path: str) -> FileResponse:
        requested_file = FRONTEND_DIST / full_path
        if full_path and requested_file.is_file():
            return FileResponse(requested_file)
        return FileResponse(FRONTEND_DIST / "index.html")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
